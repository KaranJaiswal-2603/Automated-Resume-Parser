
import re
from pathlib import Path

import pdfplumber
import spacy
from docx import Document


try:
    nlp = spacy.load("en_core_web_sm")
except OSError:
    nlp = None


KNOWN_SKILLS = [
    "Python", "Java", "C++", "C", "JavaScript", "TypeScript",
    "HTML", "CSS", "React", "Node.js", "Express.js", "Flask",
    "Django", "Spring Boot", "SQL", "PostgreSQL", "MySQL",
    "SQLite", "MongoDB", "Git", "GitHub", "Machine Learning",
    "Deep Learning", "Data Analysis", "Data Structures",
    "Algorithms", "Object-Oriented Programming", "OOP", "DBMS",
    "Operating Systems", "Pandas", "NumPy", "Matplotlib",
    "Scikit-learn", "TensorFlow", "PyTorch", "Docker", "AWS",
    "REST API", "spaCy", "Data Preprocessing", "Model Building",
    "MySQL Workbench",
]

EDUCATION_HEADINGS = {
    "education",
    "academic background",
    "educational qualifications",
    "educational qualification",
    "qualification",
    "qualifications",
}

EXPERIENCE_HEADINGS = {
    "experience",
    "work experience",
    "professional experience",
    "employment history",
    "work history",
    "internships",
    "internship experience",
    "professional background",
}

OTHER_SECTION_HEADINGS = {
    "education",
    "academic background",
    "educational qualifications",
    "technical skills",
    "technical skills and interests",
    "skills",
    "projects",
    "personal projects",
    "certifications",
    "achievements",
    "responsibilities",
    "soft skills",
    "positions of responsibility",
    "interests",
    "hobbies",
    "references",
    "declaration",
}


def normalize_heading(line):
    """Normalize a section heading."""
    return re.sub(
        r"\s+",
        " ",
        re.sub(r"[^a-z ]", "", line.lower()),
    ).strip()


def extract_text(file_path):
    """Extract readable text from a PDF or DOCX resume."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Resume file not found: {path}")

    extension = path.suffix.lower()

    if extension == ".pdf":
        text_parts = []

        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

        return "\n".join(text_parts).strip()

    if extension == ".docx":
        document = Document(path)

        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Many resumes put qualifications and contact details in tables.
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                )
                if row_text:
                    paragraphs.append(row_text)

        return "\n".join(paragraphs).strip()

    raise ValueError("Unsupported file format. Use PDF or DOCX.")


def extract_email(text):
    """Extract the first email address."""
    pattern = (
        r"\b[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"
    )

    match = re.search(pattern, text)
    return match.group(0) if match else None


def extract_phone(text):
    """Find a likely phone number."""
    pattern = (
        r"(?<!\w)"
        r"(?:\+?\d{1,3}[-.\s]?)?"
        r"(?:\(?\d{3,5}\)?[-.\s]?)?"
        r"\d{5,10}"
        r"(?!\w)"
    )

    for match in re.finditer(pattern, text):
        number = match.group(0).strip()
        digits = re.sub(r"\D", "", number)

        if len(digits) in range(10, 14):
            return number

    return None


def extract_name(text):
    """Identify a likely candidate name near the start of a resume."""
    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    if not lines:
        return None

    excluded = {skill.lower() for skill in KNOWN_SKILLS}
    excluded.update({
        "resume", "curriculum vitae", "education",
        "technical skills", "technical skills and interests",
        "skills", "experience", "work experience", "projects",
        "contact", "profile", "summary", "professional summary",
        "certifications", "achievements", "internships",
    })

    name_pattern = (
        r"[A-Za-z][A-Za-z.'-]*"
        r"(?:\s+[A-Za-z][A-Za-z.'-]*){1,3}"
    )

    for line in lines[:8]:
        cleaned = re.sub(r"\s+", " ", line).strip()
        normalized = cleaned.lower().strip(" :-|")

        if normalized in excluded:
            continue
        if "@" in cleaned or "://" in cleaned:
            continue
        if re.search(r"\d", cleaned):
            continue
        if len(cleaned.split()) > 4:
            continue

        if nlp is not None:
            doc = nlp(cleaned)

            for entity in doc.ents:
                if entity.label_ == "PERSON":
                    name = entity.text.strip()

                    if (
                        name.lower() not in excluded
                        and re.fullmatch(name_pattern, name)
                    ):
                        return name

    # Heuristic fallback when spaCy does not identify the name.
    for line in lines[:5]:
        cleaned = re.sub(r"\s+", " ", line).strip()

        if "@" in cleaned or "://" in cleaned:
            continue
        if re.search(r"\d", cleaned):
            continue
        if len(cleaned.split()) not in (2, 3):
            continue
        if any(word.lower() in excluded for word in cleaned.split()):
            continue

        if re.fullmatch(
            r"[A-Za-z][A-Za-z.'-]*"
            r"(?:\s+[A-Za-z][A-Za-z.'-]*){1,2}",
            cleaned,
        ):
            return cleaned

    return None


def extract_skills(text):
    """Find known technical skills using case-insensitive matching."""
    found_skills = []
    text_lower = text.lower()

    for skill in KNOWN_SKILLS:
        pattern = (
            r"(?<![a-z0-9])"
            + re.escape(skill.lower())
            + r"(?![a-z0-9])"
        )

        if re.search(pattern, text_lower):
            found_skills.append(skill)

    return found_skills


def extract_education(text):
    """Extract lines inside a recognized education section."""
    education_lines = []
    in_section = False

    for line in text.splitlines():
        cleaned = line.strip()

        if not cleaned:
            continue

        heading = normalize_heading(cleaned)

        if heading in EDUCATION_HEADINGS:
            in_section = True
            continue

        if in_section:
            if (
                heading in EXPERIENCE_HEADINGS
                or heading in OTHER_SECTION_HEADINGS
            ):
                break

            if "skill" in heading and len(heading.split()) <= 5:
                break

            education_lines.append(cleaned)

    return education_lines


def extract_experience(text):
    """
    Extract lines under recognized experience or internship headings.

    This is section-based extraction, not full employment-history
    classification. Review the results for unusual resume layouts.
    """
    experience_lines = []
    in_section = False

    for line in text.splitlines():
        cleaned = re.sub(r"\s+", " ", line).strip()

        if not cleaned:
            continue

        heading = normalize_heading(cleaned)

        if heading in EXPERIENCE_HEADINGS:
            in_section = True
            continue

        if in_section:
            if (
                heading in OTHER_SECTION_HEADINGS
                or heading in EDUCATION_HEADINGS
            ):
                break

            experience_lines.append(cleaned)

    result = []
    previous = None

    for line in experience_lines:
        if line != previous:
            result.append(line)
        previous = line

    return result


def parse_resume(file_path):
    """Parse the supported resume fields into a dictionary."""
    text = extract_text(file_path)

    if not text:
        raise ValueError(
            "No readable text found. The PDF may be scanned or image-based."
        )

    return {
        "full_name": extract_name(text),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "skills": extract_skills(text),
        "education": extract_education(text),
        "experience": extract_experience(text),
    }


if __name__ == "__main__":
    resume_path = input(
        "Enter the full path to a PDF or DOCX resume: "
    ).strip().strip('"')

    try:
        candidate = parse_resume(resume_path)

        print("\n--- EXTRACTED CANDIDATE DETAILS ---")

        for field, value in candidate.items():
            print(f"\n{field}:")

            if isinstance(value, list):
                if value:
                    for item in value:
                        print(f"  - {item}")
                else:
                    print("  None found")
            else:
                print(f"  {value or 'Not found'}")

    except Exception as error:
        print("Parsing failed:", repr(error))
